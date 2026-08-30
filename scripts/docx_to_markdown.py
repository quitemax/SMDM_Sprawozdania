"""Konwersja protokołów z input/historical_data/reports/ (.docx, wyjątkowo .pdf)
na Markdown (input/historical_data/reports_md/).

Zachowuje kolejność akapitów i tabel z dokumentu Word, pogrubienia/kursywę
w tekście oraz listy numerowane. Nie odtwarza dokładnego formatowania
wizualnego (wyśrodkowanie, czcionki) — celem jest wierny tekst jako
kontekst historyczny przy generowaniu sprawozdań (patrz Etap 4 w
docs/ROADMAP.md), nie odwzorowanie wyglądu dokumentu.

Jeśli dla tego samego protokołu istnieje zarówno .docx jak i .pdf
(typowo eksport tego samego dokumentu), .docx ma pierwszeństwo — czytany
bezpośrednio, bez OCR, więc dokładniejszy. Plik .pdf używany jest tylko
gdy nie ma odpowiadającego mu .docx (wtedy przez OCR, patrz
scripts/pdf_to_markdown.py — dotyczy to samo ograniczenie dokładności).

Znane ograniczenia:
- pogrubienie/kursywa wykrywane per fragment (run) — formatowanie
  odziedziczone po stylu akapitu (a nie ustawione wprost) jest pomijane;
- listy numerowane są renumerowane sekwencyjnie od 1 przy każdej nowej
  liście (numId) — nie odczytuje rzeczywistego formatu numeracji z Worda
  (np. a/b/c czy i/ii/iii), zawsze cyframi.
"""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from config import load_config
from pdf_to_markdown import convert_pdf, resolve_tesseract_cmd

MD_ESCAPE_CHARS = "*_`"


def escape_md(text: str) -> str:
    for ch in MD_ESCAPE_CHARS:
        text = text.replace(ch, "\\" + ch)
    return text


def runs_to_markdown(paragraph: Paragraph) -> str:
    # Sąsiednie fragmenty (runs) o tym samym formatowaniu są łączone przed
    # owinięciem w znaczniki Markdown — Word często dzieli jeden pogrubiony
    # fragment tekstu na kilka runs (np. osobno cyfry, osobno kropki), co
    # bez tego dawałoby np. "**Ad. ****2**" zamiast "**Ad. 2**".
    groups: list[tuple[bool, bool, str]] = []
    for run in paragraph.runs:
        if not run.text:
            continue
        style = (bool(run.bold), bool(run.italic))
        if groups and groups[-1][:2] == style:
            groups[-1] = (style[0], style[1], groups[-1][2] + run.text)
        else:
            groups.append((style[0], style[1], run.text))

    parts = []
    for bold, italic, text in groups:
        text = escape_md(text)
        if bold and italic:
            text = f"***{text}***"
        elif bold:
            text = f"**{text}**"
        elif italic:
            text = f"*{text}*"
        parts.append(text)
    return "".join(parts)


def iter_block_items(document: Document):
    """Iteruje akapity i tabele w kolejności występowania w dokumencie
    (document.paragraphs / .tables oddzielnie tracą tę kolejność)."""
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def get_numbering(paragraph: Paragraph) -> tuple[int, int] | None:
    """Zwraca (numId, ilvl) jeśli akapit należy do listy numerowanej, inaczej None."""
    p_pr = paragraph._p.pPr
    if p_pr is None or p_pr.numPr is None or p_pr.numPr.numId is None:
        return None
    ilvl = p_pr.numPr.ilvl
    return (p_pr.numPr.numId.val, ilvl.val if ilvl is not None else 0)


def table_to_markdown(table: Table) -> list[str]:
    rows = [[escape_md(cell.text.strip()).replace("\n", "<br>") for cell in row.cells] for row in table.rows]
    if not rows:
        return []
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * len(rows[0])) + " |"]
    lines += ["| " + " | ".join(row) + " |" for row in rows[1:]]
    return lines


def convert_docx(docx_path: Path, output_path: Path) -> None:
    document = Document(docx_path)
    lines: list[str] = [f"# {docx_path.stem}", ""]

    list_counters: dict[int, int] = {}
    active_num_id: int | None = None

    for block in iter_block_items(document):
        if isinstance(block, Table):
            lines += table_to_markdown(block) + [""]
            active_num_id = None
            continue

        text = runs_to_markdown(block).strip()
        numbering = get_numbering(block)

        if not text:
            lines.append("")
            active_num_id = None
            continue

        if numbering is not None:
            num_id, ilvl = numbering
            if num_id != active_num_id:
                list_counters[num_id] = 0
                active_num_id = num_id
            list_counters[num_id] += 1
            lines.append(f"{'  ' * ilvl}{list_counters[num_id]}. {text}")
        else:
            active_num_id = None
            lines.append(text)

    content = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip() + "\n"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(content, encoding="utf-8")


def find_reports(input_root: Path) -> list[Path]:
    """Zwraca listę plików do konwersji: wszystkie .docx, plus .pdf tylko
    gdy brak odpowiadającego mu .docx (patrz docstring modułu)."""
    docx_paths = sorted(input_root.rglob("*.docx"))
    docx_stems = {p.with_suffix("") for p in docx_paths}
    pdf_paths = [p for p in input_root.rglob("*.pdf") if p.with_suffix("") not in docx_stems]
    return sorted(docx_paths + pdf_paths)


def main() -> None:
    config = load_config()
    paths_cfg = config.get("paths", {})

    parser = argparse.ArgumentParser(
        description="Konwersja protokołów (.docx, wyjątkowo .pdf) z input/historical_data/reports/ na Markdown."
    )
    parser.add_argument(
        "--input-root",
        type=Path,
        default=Path(paths_cfg.get("historical_reports", "input/historical_data/reports")),
        help="Katalog ze sprawozdaniami (domyślnie z config/config.yaml).",
    )
    parser.add_argument(
        "--output-root",
        type=Path,
        default=Path(paths_cfg.get("historical_reports_md", "input/historical_data/reports_md")),
        help="Katalog wynikowy na pliki .md (domyślnie z config/config.yaml).",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Przetwórz ponownie nawet pliki, dla których .md jest nowszy niż źródłowy dokument.",
    )
    args = parser.parse_args()

    report_paths = find_reports(args.input_root)
    print(f"Znaleziono {len(report_paths)} plików do konwersji w {args.input_root}")

    tesseract_ready = False
    converted = 0
    skipped = 0
    for report_path in report_paths:
        relative = report_path.relative_to(args.input_root)
        output_path = (args.output_root / relative).with_suffix(".md")

        if not args.force and output_path.exists() and output_path.stat().st_mtime >= report_path.stat().st_mtime:
            skipped += 1
            continue

        print(f"{relative}")
        if report_path.suffix.lower() == ".docx":
            convert_docx(report_path, output_path)
        else:
            if not tesseract_ready:
                import pytesseract

                pytesseract.pytesseract.tesseract_cmd = resolve_tesseract_cmd()
                tesseract_ready = True
            print("  (brak odpowiadającego .docx — konwersja z PDF przez OCR, patrz scripts/pdf_to_markdown.py)")
            convert_pdf(report_path, output_path)
        converted += 1

    print(f"\nPrzekonwertowano: {converted}, pominięto (aktualne): {skipped}")


if __name__ == "__main__":
    main()
